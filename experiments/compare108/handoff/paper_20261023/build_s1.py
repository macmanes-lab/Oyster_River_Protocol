"""Build Supplementary Data S1 (ORP_4.0.0_lineage_results.xlsx) from the 2026-10-23 handoff csvs.
Values, not formulas, as in the previous package. Run: python build_s1.py <handoff_dir> <old_workbook> <out.xlsx>"""
import sys, json
import numpy as np, pandas as pd
from scipy.stats import wilcoxon, false_discovery_control
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

hd, oldwb, out = sys.argv[1:4]
o = pd.read_csv(f"{hd}/orp_vs_original.csv")
c = pd.read_csv(f"{hd}/orp_vs_components.csv")
V = "ORP 4.0.0"
NAMES = {"trinity": "Trinity", "spadesauto": "rnaSPAdes auto-k", "spadeshigh": "rnaSPAdes high-k", "transabyss": "Trans-ABySS"}

for s in ["orp", "original"]:
    o[f"good_pct_{s}"] = o[f"p_good_mapping_{s}"] * 100
    o[f"mapped_pct_{s}"] = o[f"p_fragments_mapped_{s}"] * 100
    o[f"contigs_k_{s}"] = o[f"contigs_{s}"] / 1000
c["good_pct"] = c.p_good_mapping * 100

# metric key, label, decimals, direction (+1 higher is better, -1 lower is better, 0 neither)
T1 = [("contigs_k", "Contigs (thousands)", 1, 0), ("mapped_pct", "Reads mapped (%)", 1, 1),
      ("good_pct", "Good mappings (%)", 1, 1), ("pytransrate_score", "pyTransRate score", 2, 1),
      ("pytransrate_optimal", "pyTransRate optimal score", 2, 1), ("unique_genes", "Unique Swiss-Prot genes", 0, 1),
      ("busco_complete", "BUSCO complete (%)", 1, 1), ("busco_single", "BUSCO single-copy (%)", 1, 1),
      ("busco_duplicated", "BUSCO duplicated (%)", 1, -1), ("busco_fragmented", "BUSCO fragmented (%)", 1, -1),
      ("busco_missing", "BUSCO missing (%)", 1, -1)]

def f(x, d):
    return f"{x:,.{d}f}"

def miq(x, d):
    q1, m, q3 = np.percentile(x, [25, 50, 75])
    return f"{f(m, d)} ({f(q1, d)}–{f(q3, d)})"

def qfmt(q):
    return "<0.001" if q < 0.001 else f"{q:.3f}"

def table1(d, prevlab):
    rows = []
    for k, lab, dec, _ in T1:
        a, b = d[f"{k}_original"].values, d[f"{k}_orp"].values
        rows.append({"Metric": lab, prevlab: miq(a, dec), V: miq(b, dec), "P": wilcoxon(b, a).pvalue,
                     "ORP higher in": f"{(b > a).sum()}/{len(d)}", "ORP lower in": f"{(b < a).sum()}/{len(d)}",
                     "Median difference (ORP − previous)": float(np.median(b - a))})
    t = pd.DataFrame(rows)
    t["q (BH, 11 metrics)"] = false_discovery_control(t.P.values)
    t["q"] = t["q (BH, 11 metrics)"].map(qfmt)
    return t

tsa, old = o[o.original_source == "TSA"], o[o.original_source != "TSA"]
t1a, t1b = table1(tsa, "TSA deposited"), table1(old, "ORP 2018")

# ---- Datasets sheet
DS = [("busco_complete", "BUSCO complete (%)"), ("busco_single", "BUSCO single-copy (%)"),
      ("busco_duplicated", "BUSCO duplicated (%)"), ("busco_fragmented", "BUSCO fragmented (%)"),
      ("busco_missing", "BUSCO missing (%)"), ("unique_genes", "Unique Swiss-Prot genes"),
      ("pytransrate_score", "pyTransRate score"), ("pytransrate_optimal", "pyTransRate optimal score"),
      ("good_pct", "Good mappings (%)"), ("mapped_pct", "Reads mapped (%)"), ("contigs", "Contigs")]
ds = pd.DataFrame({"Pair": o.pair_id, "SRA run": o.run, "Organism": o.organism, "Group": o.group,
                   "BUSCO lineage": o.busco_dataset, "BUSCO genes": o.busco_n, "Previous assembly": o.original_source,
                   "ORP build": o.orp_version.map({"4.1.0": "4.0.0 release candidate (e88d037)", "4.0.0": "4.0.0 tag (b9459aa)"}),
                   "Read pairs": o.read_pairs, "--min-cov applied": o.min_cov_applied})
for k, lab in DS:
    ds[f"{lab}: previous"] = o[f"{k}_original"]
    ds[f"{lab}: {V}"] = o[f"{k}_orp"]
    ds[f"{lab}: ORP − previous"] = o[f"{k}_orp"] - o[f"{k}_original"]
ds[f"Proper pairs (%): {V}"] = o.proper_pairs_pct_orp
ds["Wall time (h)"] = o.orp_hours
ds["Peak memory (GB)"] = o.orp_peak_mem_gb
gorder = ["mammalia", "vertebrata", "insecta", "embryophyta", "echinodermata", "other"]
ds["_g"] = ds.Group.map({g: i for i, g in enumerate(gorder)})
ds = ds.sort_values(["_g", "Previous assembly", "Organism"]).drop(columns="_g")

# ---- Groups
o["gain"] = o.busco_complete_orp - o.busco_complete_original
o["ggain"] = o.unique_genes_orp - o.unique_genes_original
grp = []
for g in gorder:
    d = o[o.group == g]
    grp.append({"Group": g, "Datasets": len(d), "TSA pairs": int((d.original_source == "TSA").sum()),
                "ORP 2018 pairs": int((d.original_source != "TSA").sum()),
                "Median BUSCO complete, previous (%)": d.busco_complete_original.median(),
                "Median BUSCO complete, ORP (%)": d.busco_complete_orp.median(), "Median gain (points)": d.gain.median(),
                "Median BUSCO duplicated, previous (%)": d.busco_duplicated_original.median(),
                "Median BUSCO duplicated, ORP (%)": d.busco_duplicated_orp.median(),
                "Duplication fell in": int((d.busco_duplicated_orp < d.busco_duplicated_original).sum()),
                "Median unique genes gain": d.ggain.median(),
                "Median pyTransRate score, previous": d.pytransrate_score_original.median(),
                "Median pyTransRate score, ORP": d.pytransrate_score_orp.median()})
groups = pd.DataFrame(grp)

# ---- Components
CM = [("busco_complete", "BUSCO complete (%)", 1), ("busco_single", "BUSCO single-copy (%)", 1),
      ("busco_duplicated", "BUSCO duplicated (%)", -1), ("busco_fragmented", "BUSCO fragmented (%)", -1),
      ("busco_missing", "BUSCO missing (%)", -1), ("unique_genes", "Unique Swiss-Prot genes", 1),
      ("pytransrate_score", "pyTransRate score", 1), ("pytransrate_optimal", "pyTransRate optimal score", 1),
      ("good_pct", "Good mappings (%)", 1), ("proper_pairs_pct", "Proper pairs (%)", 1), ("contigs", "Contigs", -1)]
W = {k: c.pivot(index="pair_id", columns="assembly", values=k) for k, _, _ in CM}
va = []
for a, an in NAMES.items():
    rr = []
    for k, lab, d in CM:
        x, y = W[k]["ORP"], W[k][a]
        diff = x - y
        rr.append({"Assembler": an, "Metric": lab, "n": len(x), "Median, assembler": y.median(), f"Median, {V}": x.median(),
                   "Median difference (ORP − assembler)": diff.median(), "Mean difference (ORP − assembler)": diff.mean(),
                   "ORP better in": int((d * diff > 0).sum()), "ORP worse in": int((d * diff < 0).sum()),
                   "Tie": int((diff == 0).sum()), "Wilcoxon P": wilcoxon(x, y).pvalue})
    q = false_discovery_control([r["Wilcoxon P"] for r in rr])
    for r, qq in zip(rr, q):
        r["q (BH within assembler)"] = qq
    va += rr
vsa = pd.DataFrame(va)

bs = []
for k, lab, d in CM:
    if k == "contigs":
        continue
    w = W[k]; comp = w[list(NAMES)]
    best = comp.max(axis=1) if d == 1 else comp.min(axis=1)
    who = (comp.idxmax(axis=1) if d == 1 else comp.idxmin(axis=1)).map(NAMES)
    diff = w["ORP"] - best
    bs.append({"Metric": lab, "Runs": len(w), "ORP better than best single": int((d * diff > 0).sum()),
               "Tie": int((diff == 0).sum()), "ORP worse than best single": int((d * diff < 0).sum()),
               "Median difference (ORP − best)": diff.median(),
               "Best single assembler, runs": "; ".join(f"{n} {v}" for n, v in who.value_counts().items())})
best = pd.DataFrame(bs)

meta = o.set_index("pair_id")[["organism", "group", "run"]]
comps = c.merge(meta, left_on="pair_id", right_index=True, suffixes=("", "_m"))
comps["Assembly"] = comps.assembly.map({**NAMES, "ORP": V})
comps["_a"] = comps.assembly.map({"ORP": 0, "trinity": 1, "spadesauto": 2, "spadeshigh": 3, "transabyss": 4})
comps = comps.sort_values(["group", "organism", "pair_id", "_a"])
components = pd.DataFrame({"Group": comps.group, "Organism": comps.organism, "Pair": comps.pair_id, "SRA run": comps.run,
                           "Assembly": comps.Assembly, **{lab: comps[k] for k, lab, _ in CM}})

# ---- RunTime: wall time and peak memory for 59; per-step times and CPU-hours from the previous package (first 50)
oldrt = pd.read_excel(oldwb, sheet_name="RunTime").set_index("Pair")
rt = pd.DataFrame({"Pair": o.pair_id, "Organism": o.organism, "Group": o.group, "Read pairs": o.read_pairs,
                   "--min-cov applied": o.min_cov_applied, "Wall time (h)": o.orp_hours, "Peak memory (GB)": o.orp_peak_mem_gb})
extra = [x for x in oldrt.columns if x.endswith("(h)") and x != "Wall time (h)"] + ["CPU-hours", "Node"]
rt = rt.join(oldrt[extra], on="Pair").sort_values("Wall time (h)")

# ---- Settings (dev30 vs 4.0.0): all 59, from settings_compare.py (copy its
# settings_dev30_vs_400.csv from lineage_410/ into the handoff directory)
sett = pd.read_csv(f"{hd}/settings_dev30_vs_400.csv")

# ---- headline numbers
h = o.orp_hours
summary = pd.DataFrame([
    ("Datasets", f"59 (44 TSA-paired, 15 paired with the ORP 2018 assembly); 57 species; six groups"),
    ("ORP vs TSA: BUSCO complete", f"median {f(tsa.busco_complete_original.median(),1)}% → {f(tsa.busco_complete_orp.median(),1)}%; higher in 44/44"),
    ("ORP vs TSA: unique genes", f"median {f(tsa.unique_genes_original.median(),0)} → {f(tsa.unique_genes_orp.median(),0)}; higher in 44/44"),
    ("ORP vs TSA: BUSCO duplicated", f"median {f(tsa.busco_duplicated_original.median(),1)}% → {f(tsa.busco_duplicated_orp.median(),1)}%"),
    ("ORP vs TSA: pyTransRate score", f"median {f(tsa.pytransrate_score_original.median(),2)} → {f(tsa.pytransrate_score_orp.median(),2)}"),
    ("ORP vs TSA: good mappings", f"median {f(tsa.good_pct_original.median(),1)}% → {f(tsa.good_pct_orp.median(),1)}%; higher in 44/44"),
    ("ORP vs ORP 2018: BUSCO complete", f"median {f(old.busco_complete_original.median(),1)}% → {f(old.busco_complete_orp.median(),1)}%; higher in 15/15"),
    ("ORP vs ORP 2018: BUSCO duplicated", f"median {f(old.busco_duplicated_original.median(),1)}% → {f(old.busco_duplicated_orp.median(),1)}%; lower in 14/15"),
    ("ORP vs ORP 2018: contigs", f"median {f(old.contigs_original.median()/1000,0)}k → {f(old.contigs_orp.median()/1000,0)}k; fewer in 15/15"),
    ("Run time (ORP 4.0.0, 32 cores)", f"median {h.median():.1f} h wall (IQR {h.quantile(.25):.1f}–{h.quantile(.75):.1f}; range {h.min():.1f}–{h.max():.1f}); median peak memory {o.orp_peak_mem_gb.median():.0f} GB"),
], columns=["Result", "Value"])

readme = pd.DataFrame([
    ("Title", "ORP 4.0.0 lineage comparison: Supplementary Data S1"),
    ("Built", "2026-10-23 from experiments/compare108/handoff (orp_vs_original.csv, orp_vs_components.csv) by build_s1.py"),
    ("Datasets", "59 paired-end RNA-seq libraries: 44 paired with the deposited TSA assembly, 15 with the ORP 2018 orthomerged assembly (Zenodo doi:10.5281/zenodo.1320141)."),
    ("ORP build", "50 runs on the 4.0.0 release candidate (commit e88d037, then labelled 4.1.0); 9 ORP 2018 libraries on the tagged 4.0.0 release (b9459aa). Between the two commits only resume handling and a startup banner changed."),
    ("BUSCO", "v6.1.0, transcriptome mode, offline, ORP's config.ini; one lineage dataset per group (column BUSCO lineage). mammalia, vertebrata and insecta are OrthoDB 12.1 (odb12); the rest OrthoDB 12.2."),
    ("pyTransRate", "2.2.2 for every assembly (ORP, previous and the four components), with the trimmed, error-corrected reads from that library's ORP run."),
    ("Unique genes", "Distinct Swiss-Prot gene names among DIAMOND blastx hits (E ≤ 1e-8, --top 0.1), counted as ORP counts them."),
    ("Statistics", "Two-sided Wilcoxon signed-rank tests on paired values; Benjamini-Hochberg q within each Table 1 panel (11 metrics) or within each assembler (11 metrics)."),
    ("Values", "All cells are values computed in Python."),
    ("Sheets", "Summary; Table1a; Table1b; Datasets; Groups; ORP_vs_assembler; ORP_vs_best_single; Components; Settings_dev30_vs_400 (Section 4.4, all 59 datasets: ORP 4.0.0 against a run with read normalization on and a fixed --min-cov 2; column 'dev30-settings run' says whether that run was the dev30 build (first 50) or ORP 4.0.0 with those two settings (last nine)); RunTime (per-step hours and CPU-hours for the first 50 only)."),
    ("Known caveats", "tsa_GAGE reads are Apiospora malaysiana (ascomycete) but both assemblies were scored against agaricales_odb12.2. TSA assemblies are curated (no contigs < 200 bp). tsa_GAQK excluded (R2 reads 62 bp)."),
], columns=["Item", "Detail"])

sheets = {"README": readme, "Summary": summary,
          "Table1a": t1a.drop(columns=["q (BH, 11 metrics)"]).assign(**{"q (BH, 11 metrics)": t1a["q (BH, 11 metrics)"]}),
          "Table1b": t1b.drop(columns=["q (BH, 11 metrics)"]).assign(**{"q (BH, 11 metrics)": t1b["q (BH, 11 metrics)"]}),
          "Datasets": ds, "Groups": groups, "ORP_vs_assembler": vsa, "ORP_vs_best_single": best,
          "Components": components, "Settings_dev30_vs_400": sett, "RunTime": rt}
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    for n, d in sheets.items():
        d.to_excel(xw, sheet_name=n, index=False)

wb = load_workbook(out)
for ws in wb:
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
            if cell.row == 1:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col[1:50]) if ws.max_row > 1 else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60 if ws.title in ("README", "Summary") else 28)
wb.save(out)

json.dump({"t1a": t1a.drop(columns="P").to_dict("records"), "t1b": t1b.drop(columns="P").to_dict("records"),
           "groups": groups.to_dict("records"), "best": best.to_dict("records")},
          open(out.replace(".xlsx", "_numbers.json"), "w"), indent=1, default=float)
print("ok")
